"""
front_matter.py
===============
Builds all front matter sections for the B.Tech project report:
- Cover Page
- Bonafide Certificate
- Declaration
- Abstract (covering all 19 required points)
- Acknowledgement
- Table of Contents
- List of Abbreviations
- List of Figures
- List of Tables
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from .styles import (
    FONT_NAME, COLOR_BLACK, COLOR_NAVY, COLOR_DARK_GRAY,
    add_prelim_heading, add_body_p, add_table_data,
    set_cell_shading, set_cell_margins
)

def add_cover_page(doc):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(20)
    p_title.paragraph_format.line_spacing = 1.2
    
    r_title = p_title.add_run(
        "AN EXPLAINABLE HYBRID FEATURE-FUSION FRAMEWORK FOR RICE QUALITY "
        "AND DEFECT CLASSIFICATION USING EFFICIENTNET-B0 AND XGBOOST"
    )
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_BLACK

    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_proj.paragraph_format.space_before = Pt(12)
    p_proj.paragraph_format.space_after = Pt(18)
    r_proj = p_proj.add_run("27PR27 — PROJECT WORK II")
    r_proj.font.name = FONT_NAME
    r_proj.font.size = Pt(13)
    r_proj.font.bold = True
    r_proj.font.color.rgb = COLOR_BLACK
    
    p_rep = doc.add_paragraph()
    p_rep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rep.paragraph_format.space_before = Pt(0)
    p_rep.paragraph_format.space_after = Pt(20)
    r_rep = p_rep.add_run("A PROJECT REPORT\n\nSubmitted by")
    r_rep.font.name = FONT_NAME
    r_rep.font.size = Pt(12)
    r_rep.font.bold = True
    r_rep.font.color.rgb = COLOR_BLACK

    p_stud = doc.add_paragraph()
    p_stud.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_stud.paragraph_format.space_before = Pt(6)
    p_stud.paragraph_format.space_after = Pt(24)
    p_stud.paragraph_format.line_spacing = 1.3
    
    r_s1 = p_stud.add_run("DHARSHAN B\n(Roll No: 23ITR030)\n\n")
    r_s1.font.name = FONT_NAME
    r_s1.font.size = Pt(12)
    r_s1.font.bold = True
    
    r_s2 = p_stud.add_run("DINESH G L\n(Roll No: 23ITR039)")
    r_s2.font.name = FONT_NAME
    r_s2.font.size = Pt(12)
    r_s2.font.bold = True

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_deg.paragraph_format.space_before = Pt(12)
    p_deg.paragraph_format.space_after = Pt(16)
    p_deg.paragraph_format.line_spacing = 1.15
    r_deg = p_deg.add_run(
        "in partial fulfillment of the requirements\n"
        "for the award of the degree of\n\n"
        "BACHELOR OF TECHNOLOGY\nIN\nINFORMATION TECHNOLOGY"
    )
    r_deg.font.name = FONT_NAME
    r_deg.font.size = Pt(12)
    r_deg.font.bold = True

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_before = Pt(20)
    p_dept.paragraph_format.space_after = Pt(0)
    p_dept.paragraph_format.line_spacing = 1.2
    r_dept = p_dept.add_run(
        "DEPARTMENT OF INFORMATION TECHNOLOGY\n"
        "KONGU ENGINEERING COLLEGE\n"
        "(AUTONOMOUS)\n"
        "PERUNDURAI, ERODE – 638 060\n"
        "ACADEMIC YEAR: 2026–2027"
    )
    r_dept.font.name = FONT_NAME
    r_dept.font.size = Pt(12)
    r_dept.font.bold = True

def add_bonafide_certificate(doc):
    doc.add_page_break()
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(12)
    p_inst.paragraph_format.space_after = Pt(18)
    p_inst.paragraph_format.line_spacing = 1.15
    r_inst = p_inst.add_run(
        "DEPARTMENT OF INFORMATION TECHNOLOGY\n"
        "KONGU ENGINEERING COLLEGE\n"
        "(AUTONOMOUS)\n"
        "PERUNDURAI, ERODE – 638 060\n"
        "OCTOBER 2026\n\n"
        "BONAFIDE CERTIFICATE"
    )
    r_inst.font.name = FONT_NAME
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True

    add_body_p(
        doc,
        "This is to certify that the project report titled \"AN EXPLAINABLE HYBRID FEATURE-FUSION "
        "FRAMEWORK FOR RICE QUALITY AND DEFECT CLASSIFICATION USING EFFICIENTNET-B0 AND XGBOOST\" "
        "is the bonafide record of project work done by DHARSHAN B (Roll No: 23ITR030) and DINESH G L "
        "(Roll No: 23ITR039) in partial fulfillment of the requirements for the award of the Degree of "
        "Bachelor of Technology in Information Technology of Anna University, Chennai during the "
        "academic year 2026–2027.",
        space_after=28, line_spacing=1.5
    )

    # Supervisor and HOD signature blocks
    tbl_sigs = doc.add_table(rows=1, cols=2)
    tbl_sigs.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sigs.autofit = False
    
    cell_sup = tbl_sigs.cell(0, 0)
    cell_sup.width = Inches(3.2)
    p_sup = cell_sup.paragraphs[0]
    p_sup.paragraph_format.line_spacing = 1.15
    r_sup = p_sup.add_run(
        "SUPERVISOR\n"
        "(Ms. SRIPRIYA S)\n"
        "Assistant Professor\n"
        "Department of Information Technology\n"
        "Kongu Engineering College (Autonomous)\n"
        "Perundurai, Erode – 638 060"
    )
    r_sup.font.name = FONT_NAME
    r_sup.font.size = Pt(11)
    r_sup.font.bold = True

    cell_hod = tbl_sigs.cell(0, 1)
    cell_hod.width = Inches(3.2)
    p_hod = cell_hod.paragraphs[0]
    p_hod.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_hod.paragraph_format.line_spacing = 1.15
    r_hod = p_hod.add_run(
        "HEAD OF THE DEPARTMENT\n"
        "(Dr. S. ANANDAMURUGAN)\n"
        "Professor and Head\n"
        "Department of Information Technology\n"
        "Kongu Engineering College (Autonomous)\n"
        "Perundurai, Erode – 638 060"
    )
    r_hod.font.name = FONT_NAME
    r_hod.font.size = Pt(11)
    r_hod.font.bold = True

    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_before = Pt(36)
    p_date.paragraph_format.space_after = Pt(16)
    r_date = p_date.add_run("Date: ")
    r_date.font.name = FONT_NAME
    r_date.font.size = Pt(11)
    r_date.font.bold = True

    p_viva = doc.add_paragraph()
    p_viva.paragraph_format.space_before = Pt(12)
    p_viva.paragraph_format.space_after = Pt(40)
    r_viva = p_viva.add_run("Submitted for the end semester viva voce examination held on ____________________")
    r_viva.font.name = FONT_NAME
    r_viva.font.size = Pt(11)

    tbl_ex = doc.add_table(rows=1, cols=2)
    tbl_ex.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ex.autofit = False
    c_in = tbl_ex.cell(0, 0)
    c_in.width = Inches(3.2)
    r_in = c_in.paragraphs[0].add_run("INTERNAL EXAMINER")
    r_in.font.name = FONT_NAME
    r_in.font.size = Pt(11)
    r_in.font.bold = True
    
    c_ex = tbl_ex.cell(0, 1)
    c_ex.width = Inches(3.2)
    p_ex = c_ex.paragraphs[0]
    p_ex.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ex = p_ex.add_run("EXTERNAL EXAMINER")
    r_ex.font.name = FONT_NAME
    r_ex.font.size = Pt(11)
    r_ex.font.bold = True

def add_declaration(doc):
    doc.add_page_break()
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(12)
    p_inst.paragraph_format.space_after = Pt(18)
    p_inst.paragraph_format.line_spacing = 1.15
    r_inst = p_inst.add_run(
        "DEPARTMENT OF INFORMATION TECHNOLOGY\n"
        "KONGU ENGINEERING COLLEGE\n"
        "(AUTONOMOUS)\n"
        "PERUNDURAI, ERODE – 638 060\n"
        "OCTOBER 2026\n\n"
        "DECLARATION"
    )
    r_inst.font.name = FONT_NAME
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True

    add_body_p(
        doc,
        "We affirm that the Project Report titled \"AN EXPLAINABLE HYBRID FEATURE-FUSION "
        "FRAMEWORK FOR RICE QUALITY AND DEFECT CLASSIFICATION USING EFFICIENTNET-B0 AND XGBOOST\" "
        "being submitted in partial fulfillment of the requirements for the award of Bachelor of "
        "Technology in Information Technology is the original work carried out by us. It has not formed "
        "part of any other project report or dissertation based on which a degree or award was conferred on an "
        "earlier occasion on this or any other candidate.",
        space_after=24, line_spacing=1.5
    )

    tbl_cands = doc.add_table(rows=1, cols=2)
    tbl_cands.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cands.autofit = False
    
    c1 = tbl_cands.cell(0, 0)
    c1.width = Inches(3.2)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.line_spacing = 1.15
    r1 = p1.add_run("DHARSHAN B\n(Roll No: 23ITR030)")
    r1.font.name = FONT_NAME
    r1.font.size = Pt(11)
    r1.font.bold = True

    c2 = tbl_cands.cell(0, 1)
    c2.width = Inches(3.2)
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run("DINESH G L\n(Roll No: 23ITR039)")
    r2.font.name = FONT_NAME
    r2.font.size = Pt(11)
    r2.font.bold = True

    p_dt = doc.add_paragraph()
    p_dt.paragraph_format.space_before = Pt(28)
    p_dt.paragraph_format.space_after = Pt(24)
    r_dt = p_dt.add_run("Date: ")
    r_dt.font.name = FONT_NAME
    r_dt.font.size = Pt(11)
    r_dt.font.bold = True

    add_body_p(
        doc,
        "I certify that the declaration made by the above candidates is true to the best of my knowledge.",
        space_after=36, line_spacing=1.3
    )

    p_sup_endorse = doc.add_paragraph()
    p_sup_endorse.paragraph_format.space_before = Pt(10)
    p_sup_endorse.paragraph_format.space_after = Pt(0)
    p_sup_endorse.paragraph_format.line_spacing = 1.15
    r_se = p_sup_endorse.add_run(
        "Date:\t\t\t\t\t\tName and Signature of the Supervisor with Seal\n"
        "\t\t\t\t\t\t\t(Ms. SRIPRIYA S, M.E.)"
    )
    r_se.font.name = FONT_NAME
    r_se.font.size = Pt(11)
    r_se.font.bold = True

def add_abstract_section(doc):
    doc.add_page_break()
    add_prelim_heading(doc, "ABSTRACT")

    # 1. Importance of automated rice quality and defect assessment
    # 2. Limitations of purely manual inspection
    # 3. Need for computer vision and machine learning
    # 4. OpenCV-based preprocessing and segmentation
    # 5. Extraction of handcrafted features
    # 6. Shape/morphological features
    # 7. GLCM texture features
    # 8. Colour features
    # 9. EfficientNet-B0 deep feature extraction
    # 10. Feature fusion
    # 11. 1,342-dimensional combined representation
    # 12. StandardScaler
    # 13. XGBoost classification
    # 14. Eight-class classification
    # 15. SHAP explainability
    # 16. Experimental evaluation
    # 17. Final measured performance
    # 18. Practical deployment support
    # 19. Limitations and future potential
    p1 = (
        "Rice (Oryza sativa) serves as a primary staple crop sustaining over half of the global population, "
        "making accurate, objective, and non-destructive post-harvest grain quality and defect grading essential "
        "for trade equity, food security, and consumer protection. Traditional inspection in commercial milling "
        "and procurement facilities relies heavily on manual visual evaluation by human sorters. This manual paradigm "
        "suffers from substantial operational bottlenecks, including high subjectivity, operator fatigue, inter-observer "
        "inconsistency, labor intensity, and an inability to sustain rapid inspection rates across industrial volumes. "
        "Consequently, there is an urgent industrial requirement for automated, reproducible image-based inspection "
        "frameworks leveraging computer vision and machine learning to achieve rapid and reliable defect grading."
    )
    add_body_p(doc, p1, line_spacing=1.5, space_after=10, indent=0.3)

    p2 = (
        "In this work, an explainable hybrid feature-fusion framework is proposed and systematically evaluated for "
        "multiclass rice grain categorization across eight raw classes: 0_NOR, 1_F&S, 2_SD, 3_MY, 4_AP, 5_BN, 6_UN, "
        "and 7_IM. To establish an artifact-free input representation, an OpenCV-based preprocessing and segmentation "
        "pipeline is implemented, comprising Gaussian spatial filtering (5×5 kernel), Otsu automatic thresholding, "
        "elliptical morphological closing and opening, external contour isolation, and aspect-ratio preserved isotropic "
        "resizing with centered background padding to a standardized 224×224×3 pixel format. The system constructs two "
        "parallel feature extraction branches. Branch 1 extracts 62 domain-specific handcrafted descriptors comprising "
        "14 shape and morphological metrics (including area, perimeter, aspect ratio, solidity, eccentricity, and axis "
        "lengths), 12 Gray-Level Co-occurrence Matrix (GLCM) texture descriptors (contrast, dissimilarity, homogeneity, "
        "energy, correlation, and angular second moment aggregated across spatial offsets of 1 and 2 pixels at angles "
        "0°, 45°, 90°, and 135°), and 36 colour distribution statistics (mean, standard deviation, minimum, and maximum "
        "computed across 9 chromatic channels spanning RGB, HSV, and CIELAB colour spaces). Branch 2 leverages a pretrained "
        "EfficientNet-B0 convolutional neural network with its terminal classification head excised to extract a 1,280-dimensional "
        "deep latent representation."
    )
    add_body_p(doc, p2, line_spacing=1.5, space_after=10, indent=0.3)

    p3 = (
        "The extracted handcrafted and deep representations are horizontally concatenated into a unified 1,342-dimensional "
        "fused feature vector (62 handcrafted + 1,280 deep). To eliminate data leakage and prevent numeric scale domination "
        "across heterogeneous feature modalities, a StandardScaler transformation is fitted strictly on training data and "
        "subsequently applied to validation and test splits. The standardized representations are classified using an optimized "
        "Extreme Gradient Boosting (XGBoost) model configured with n_estimators=100, max_depth=4, and learning_rate=0.1. "
        "To provide operational transparency and auditability, SHapley Additive exPlanations (SHAP) via TreeExplainer are "
        "integrated to compute exact local feature attributions and global feature-group contributions, explicitly "
        "distinguishing training split-gain from marginal test instance attribution."
    )
    add_body_p(doc, p3, line_spacing=1.5, space_after=10, indent=0.3)

    p4 = (
        "Rigorous empirical evaluation was conducted on the GrainSet Rice Dataset comprising 30,962 images partitioned "
        "into a strict zero-leakage split of 24,767 training images, 3,095 validation images, and 3,100 test images. "
        "Due to severe class imbalance—where the majority normal class (0_NOR) comprises approximately 64.5% of the data—Macro F1 "
        "serves as the primary comparative objective alongside accuracy and weighted F1. On the untouched 3,100-sample test split, "
        "the proposed hybrid framework achieves a final headline accuracy of 92.10% (92.0968%), a Macro Precision of 0.8454, "
        "a Macro Recall of 0.8940, a Macro F1 score of 0.8660, and a Weighted F1 score of 0.9244. This decisively outperforms "
        "both the classical SVM baseline operating on 62 handcrafted features (91.16% accuracy, 0.8494 Macro F1) and the deep-only "
        "XGBoost baseline operating on 1,280 EfficientNet-B0 embeddings (90.65% accuracy, 0.8399 Macro F1). Ablation analysis "
        "empirically validates that handcrafted geometry and colour provide non-redundant discriminative boundaries that deep "
        "embeddings fail to capture in isolation."
    )
    add_body_p(doc, p4, line_spacing=1.5, space_after=10, indent=0.3)

    p5 = (
        "To support practical mill-floor decision making, the framework is operationalized through an interactive Streamlit "
        "interface featuring single-grain analysis, deep visual saliency previews, prediction probability tables, a human-in-the-loop "
        "review queue for low-probability predictions, SQLite historical persistence, and a headless FastAPI REST service. "
        "While model probabilities provide useful ranking heuristics, they reflect uncalibrated softmax outputs rather than true "
        "posterior certainty, and SHAP attributions represent local feature contributions rather than proofs of model correctness. "
        "Future enhancements include probability calibration, multi-grain simultaneous segmentation, mill-floor lighting robustness, "
        "and edge-device ONNX deployment."
    )
    add_body_p(doc, p5, line_spacing=1.5, space_after=12, indent=0.3)

def add_acknowledgement_section(doc):
    doc.add_page_break()
    add_prelim_heading(doc, "ACKNOWLEDGEMENT")

    p1 = (
        "First and foremost, we express our profound gratitude and humble reverence to the Almighty for granting us "
        "the wisdom, endurance, good health, and guiding light necessary to carry out and successfully complete this project work."
    )
    add_body_p(doc, p1, line_spacing=1.5, space_after=8, indent=0.3)

    p2 = (
        "We wish to express our heartfelt gratefulness and high esteem to our beloved Correspondent, "
        "Thiru. E. R. K. KRISHNAN, M.Com., and all the distinguished members of the Kongu Vellalar Institute of "
        "Technology Trust (KVITT) for establishing exceptional academic infrastructure, world-class computational laboratories, "
        "and fostering an inspiring research environment."
    )
    add_body_p(doc, p2, line_spacing=1.5, space_after=8, indent=0.3)

    p3 = (
        "We convey our sincere and deep sense of gratitude to our respected Principal, Dr. R. PARAMESHWARAN, M.E., Ph.D., "
        "for his constant encouragement, administrative support, and for providing the institutional ecosystem to pursue "
        "meaningful applied engineering projects."
    )
    add_body_p(doc, p3, line_spacing=1.5, space_after=8, indent=0.3)

    p4 = (
        "We record our profound indebtedness and sincere thanks to Dr. S. ANANDAMURUGAN, M.E., Ph.D., Head of the Department, "
        "Department of Information Technology, for his inspiring leadership, valuable technical suggestions, and continuous "
        "administrative guidance throughout the course of our curriculum and project phases."
    )
    add_body_p(doc, p4, line_spacing=1.5, space_after=8, indent=0.3)

    p5 = (
        "We express our special thanks to our Project Coordinators, Ms. K. SRUTHI, M.E., and Ms. R. SANDHIYA, M.E., "
        "Assistant Professors, Department of Information Technology, for their periodic reviews, constructive critiques, "
        "and well-structured evaluation schedules that kept our work focused and on schedule."
    )
    add_body_p(doc, p5, line_spacing=1.5, space_after=8, indent=0.3)

    p6 = (
        "We are particularly beholden and express our deepest gratitude to our project supervisor and mentor, "
        "Ms. SRIPRIYA S, M.E., Assistant Professor, Department of Information Technology, for her invaluable guidance, "
        "patient supervision, technical scrutiny, and insightful discussions throughout the phases of data hygiene audit, "
        "feature engineering, model training, and manuscript refinement."
    )
    add_body_p(doc, p6, line_spacing=1.5, space_after=8, indent=0.3)

    p7 = (
        "We also gratefully acknowledge the Centre of Excellence in Data Science and the high-performance computing "
        "laboratories of the Department of Information Technology, Kongu Engineering College, for facilitating high-performance "
        "workstation infrastructure essential for deep feature extraction, model benchmarking, and SHAP ablation studies."
    )
    add_body_p(doc, p7, line_spacing=1.5, space_after=8, indent=0.3)

    p8 = (
        "Finally, we express our warmest love, gratitude, and indebtedness to our beloved parents, family members, and friends "
        "whose continuous sacrifices, emotional encouragement, moral support, and patience have been the ultimate foundation of "
        "our academic journey."
    )
    add_body_p(doc, p8, line_spacing=1.5, space_after=12, indent=0.3)

def add_table_of_contents(doc):
    doc.add_page_break()
    add_prelim_heading(doc, "TABLE OF CONTENTS")

    toc_headers = ["CHAPTER NO.", "TITLE", "PAGE NO."]
    toc_data = [
        ["", "ABSTRACT", "iv"],
        ["", "ACKNOWLEDGEMENT", "v"],
        ["", "LIST OF ABBREVIATIONS", "viii"],
        ["", "LIST OF FIGURES", "ix"],
        ["", "LIST OF TABLES", "x"],
        ["1", "INTRODUCTION", "1"],
        ["", "1.1 Introduction", "1"],
        ["", "1.2 Objective", "2"],
        ["", "1.3 Scope", "3"],
        ["2", "LITERATURE REVIEW", "5"],
        ["", "2.1 Thematic Literature Survey", "5"],
        ["", "2.2 Summary of Reviewed Literature (Table 2.1)", "9"],
        ["3", "PROBLEM DEFINITION", "11"],
        ["", "3.1 Existing System", "11"],
        ["", "3.2 Problem Statement", "12"],
        ["4", "SYSTEM REQUIREMENTS", "13"],
        ["", "4.1 Hardware Requirements", "13"],
        ["", "4.2 Software Requirements", "13"],
        ["", "4.3 Software Description", "14"],
        ["", "    4.3.1 Python", "14"],
        ["", "    4.3.2 OpenCV", "14"],
        ["", "    4.3.3 NumPy", "14"],
        ["", "    4.3.4 Pandas", "15"],
        ["", "    4.3.5 Scikit-learn", "15"],
        ["", "    4.3.6 PyTorch", "15"],
        ["", "    4.3.7 EfficientNet-B0", "15"],
        ["", "    4.3.8 XGBoost", "16"],
        ["", "    4.3.9 SHAP", "16"],
        ["", "    4.3.10 Streamlit", "16"],
        ["", "    4.3.11 SQLite", "17"],
        ["", "    4.3.12 FastAPI", "17"],
        ["5", "SYSTEM IMPLEMENTATION", "18"],
        ["", "5.1 Proposed System", "18"],
        ["", "    5.1.1 Image Preprocessing and Grain Segmentation", "18"],
        ["", "    5.1.2 Handcrafted Feature Extraction Pipeline", "20"],
        ["", "    5.1.3 GLCM Texture Features", "21"],
        ["", "    5.1.4 Multichannel Colour Features", "22"],
        ["", "    5.1.5 Support Vector Machine (SVM) Baseline", "23"],
        ["", "    5.1.6 StandardScaler and Data Leakage Prevention", "24"],
        ["", "    5.1.7 EfficientNet-B0 Deep Feature Backbone", "25"],
        ["", "    5.1.8 Feature Fusion Architecture", "26"],
        ["", "    5.1.9 Extreme Gradient Boosting (XGBoost) Classifier", "27"],
        ["", "    5.1.10 SHAP Explainability Framework", "29"],
        ["", "5.2 System Architecture and Workflow", "31"],
        ["6", "RESULTS AND DISCUSSION", "33"],
        ["", "6.1 Dataset Summary and Partitioning", "33"],
        ["", "6.2 Image Preprocessing Evaluation", "34"],
        ["", "6.3 Handcrafted Feature Extraction Integrity", "34"],
        ["", "6.4 Baseline Model Evaluation", "35"],
        ["", "6.5 EfficientNet-B0 + XGBoost Baseline", "35"],
        ["", "6.6 Final Hybrid Model Evaluation", "36"],
        ["", "6.7 Class-Wise Classification Dynamics", "37"],
        ["", "6.8 Confusion Matrix Analysis", "38"],
        ["", "6.9 Systematic Feature Ablation Study", "39"],
        ["", "6.10 SHAP Interpretability and Group Contributions", "41"],
        ["", "6.11 XGBoost Split-Gain vs. SHAP Attributions", "42"],
        ["", "6.12 Detailed Error and Boundary Leakage Analysis", "43"],
        ["", "6.13 Deployment Stack and Interactive Application", "44"],
        ["", "6.14 Output Probability Interpretation and Calibration Nuances", "45"],
        ["7", "CONCLUSION AND FUTURE WORK", "46"],
        ["", "7.1 Conclusion", "46"],
        ["", "7.2 Future Work", "47"],
        ["8", "APPENDICES", "48"],
        ["", "8.1 Appendix – 1 Coding", "48"],
        ["", "8.2 Appendix – 2 Output", "56"],
        ["", "REFERENCES", "63"],
        ["", "GENAI DISCLOSURE", "65"]
    ]
    add_table_data(doc, toc_headers, toc_data, col_widths=[1.2, 4.0, 1.0])

def add_list_of_abbreviations(doc):
    doc.add_page_break()
    add_prelim_heading(doc, "LIST OF ABBREVIATIONS")

    abbr_headers = ["ABBREVIATION", "EXPANSION / MEANING"]
    abbr_data = [
        ["AI", "Artificial Intelligence"],
        ["API", "Application Programming Interface"],
        ["ASM", "Angular Second Moment"],
        ["AUC", "Area Under the Receiver Operating Characteristic Curve"],
        ["CAD", "Computer-Aided Diagnosis / Detection"],
        ["CIELAB / LAB", "Commission Internationale de l'Éclairage L*a*b* Colour Space"],
        ["CNN", "Convolutional Neural Network"],
        ["CSV", "Comma-Separated Values"],
        ["DL", "Deep Learning"],
        ["GLCM", "Gray-Level Co-occurrence Matrix"],
        ["GPU", "Graphics Processing Unit"],
        ["HSV", "Hue Saturation Value Colour Space"],
        ["IDE", "Integrated Development Environment"],
        ["IEEE", "Institute of Electrical and Electronics Engineers"],
        ["KEC", "Kongu Engineering College"],
        ["MBConv", "Mobile Inverted Bottleneck Convolution"],
        ["ML", "Machine Learning"],
        ["Otsu", "Otsu Automatic Bimodal Thresholding Method"],
        ["PDF", "Portable Document Format"],
        ["RBF", "Radial Basis Function"],
        ["RGB", "Red Green Blue Colour Space"],
        ["ROC", "Receiver Operating Characteristic"],
        ["SE", "Squeeze-and-Excitation"],
        ["SHAP", "SHapley Additive exPlanations"],
        ["SQLite", "Structured Query Language Lite Embedded Relational Database"],
        ["SSD", "Solid-State Drive"],
        ["SVM / SVC", "Support Vector Machine / Support Vector Classifier"],
        ["TreeSHAP", "Tree-based SHapley Additive exPlanations Algorithm"],
        ["UI", "User Interface"],
        ["XAI", "Explainable Artificial Intelligence"],
        ["XGBoost", "Extreme Gradient Boosting"]
    ]
    add_table_data(doc, abbr_headers, abbr_data, col_widths=[1.8, 4.4])

def add_list_of_figures(doc):
    doc.add_page_break()
    add_prelim_heading(doc, "LIST OF FIGURES")

    fig_headers = ["FIGURE NO.", "FIGURE TITLE", "PAGE NO."]
    fig_data = [
        ["Figure 5.1.1", "Multi-stage Image Preprocessing and Foreground Grain Isolation Pipeline", "19"],
        ["Figure 5.1.2", "Four-Stage Preprocessing Workflow Demonstrated on Representative Grain Sample", "20"],
        ["Figure 5.2.1", "End-to-End Architectural Pipeline of the Proposed Explainable Hybrid Framework", "32"],
        ["Figure 6.2.1", "Preprocessing Verification Summary Grid across Representative Classes", "34"],
        ["Figure 6.4.1", "Benchmark Model Comparison: Handcrafted SVM vs. EfficientNet-B0 vs. Hybrid", "35"],
        ["Figure 6.8.1", "Raw Confusion Matrix of the Proposed Hybrid Model on 3,100 Test Samples", "38"],
        ["Figure 6.8.2", "Normalized Confusion Matrix Demonstrating Class-Wise Classification Dynamics", "38"],
        ["Figure 6.9.1", "Systematic Feature Ablation Study: Validation and Test Scores across 7 Configurations", "40"],
        ["Figure 6.10.1", "Top 20 Features Ranked by Mean Absolute SHAP Attribution on 300 Test Instances", "41"],
        ["Figure 6.10.2", "SHAP Beeswarm Summary Plot Displaying Feature Value Directional Impact", "41"],
        ["Figure 6.10.3", "SHAP Feature Group Attribution Share across Feature Subsets", "42"],
        ["Figure 6.11.1", "Dual Comparison: XGBoost Tree-Gain Split Importance vs. Instance-Level SHAP Share", "43"],
        ["Figure 6.12.1", "Per-Class F1-Score Comparison across Benchmark Models", "44"],
        ["Figure 8.2.1", "Original Rice Image Sample before Preprocessing", "56"],
        ["Figure 8.2.2", "Preprocessed Grayscale Denoised Image Sample", "56"],
        ["Figure 8.2.3", "Segmented Grain Foreground Binary Mask", "57"],
        ["Figure 8.2.4", "Extracted 62-Dimensional Handcrafted Feature Vector Distribution", "57"],
        ["Figure 8.2.5", "Standardized 224×224×3 Isotropically Scaled Grain Representation", "58"],
        ["Figure 8.2.6", "EfficientNet-B0 1,280-Dimensional Latent Embedding Extraction", "58"],
        ["Figure 8.2.7", "Complete System Architecture Framework Diagram", "59"],
        ["Figure 8.2.8", "Three-Model Benchmark Evaluation Bar Chart", "59"],
        ["Figure 8.2.9", "Hybrid XGBoost Final Model Test Confusion Matrix", "60"],
        ["Figure 8.2.10", "SHAP Feature Attribution Waterfall and Summary Plots", "60"],
        ["Figure 8.2.11", "Streamlit Interactive Single-Grain Inspection Dashboard", "61"],
        ["Figure 8.2.12", "Streamlit Prediction Probability Table and Agronomic Diagnostic Output", "61"],
        ["Figure 8.2.13", "Streamlit Batch Quality Analysis and CSV Export Interface", "62"],
        ["Figure 8.2.14", "Automated Commercial Rice Quality PDF Inspection Report Output", "62"]
    ]
    add_table_data(doc, fig_headers, fig_data, col_widths=[1.4, 4.0, 0.8])

def add_list_of_tables(doc):
    doc.add_page_break()
    add_prelim_heading(doc, "LIST OF TABLES")

    tbl_headers = ["TABLE NO.", "TABLE TITLE", "PAGE NO."]
    tbl_data = [
        ["Table 2.1", "Comprehensive Literature Review and Comparative Assessment Matrix", "9"],
        ["Table 3.1", "Qualitative Comparison: Existing Manual/Monolithic Systems vs. Proposed Framework", "12"],
        ["Table 4.1", "Development Workstation Hardware Configuration", "13"],
        ["Table 4.2", "Verified Software Stack and Library Versions", "14"],
        ["Table 5.1", "Complete Inventory and Mathematical Description of 62 Handcrafted Features", "22"],
        ["Table 5.2", "Verified Hyperparameter Configuration for Baseline and Final Hybrid Models", "28"],
        ["Table 6.1", "GrainSet Rice Dataset Distribution across Partitioned Zero-Leakage Splits", "33"],
        ["Table 6.2", "Benchmark Model Performance Evaluation on 3,100 Test Samples", "35"],
        ["Table 6.3", "Final Hybrid XGBoost Model Validation and Test Performance Summary", "36"],
        ["Table 6.4", "Class-Wise Classification Dynamics for the Final Hybrid XGBoost Model", "37"],
        ["Table 6.5", "Complete 8×8 Test Confusion Matrix on 3,100 Evaluation Samples", "38"],
        ["Table 6.6", "Systematic Seven-Configuration Feature Ablation Study Results", "40"],
        ["Table 6.7", "SHAP Feature Group Contribution Summary across 300 Test Instances", "42"],
        ["Table 6.8", "Comparative Breakdown: XGBoost Tree-Gain Split Importance vs. SHAP Attribution", "43"]
    ]
    add_table_data(doc, tbl_headers, tbl_data, col_widths=[1.4, 4.0, 0.8])
