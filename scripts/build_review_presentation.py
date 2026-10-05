"""
build_review_presentation.py
============================
Populates the college PPT template (kec template.pptx) with the complete,
rigorous, and presentation-ready content for PROJECT 27PR27:
"An Explainable Hybrid Feature-Fusion Framework for Rice Quality and Defect
Classification Using EfficientNet-B0 and XGBoost".

Follows the college-required sections:
1. Title of the Project
2. Objective(s)
3. Introduction
4. Literature Review (Minimum 7-10 papers, base paper as Ref [1], 2023-2026 reputed journals)
5. Summary of Literature Review
6. Problem Description / Existing Method
7. Proposed Methodology (Block Diagram, Flowchart)
8. Software / Tools and Dataset to be Used
9. Description of Each Phase / Block
10. Demo and Results (Benchmarks, Per-class, Confusion Matrix, Ablation, SHAP)
11. Screenshots (Preprocessing & Deployment)
12. References (IEEE format)
13. Thank you

Uses SIMPLE, audience-friendly English as instructed.
Preserves all verified project values and college branding.
"""

from pathlib import Path
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PROJECT_ROOT = Path("c:/Rice classifier final project")
TEMPLATE_PPTX = PROJECT_ROOT / "kec template.pptx"
OUTPUT_PPTX = PROJECT_ROOT / "27PR27_Rice_Quality_Defect_Project_Review.pptx"
FIG_DIR = PROJECT_ROOT / "paper" / "figures"
SAMPLE_DIR = PROJECT_ROOT / "results" / "preprocessing" / "intermediate_samples" / "sample_01_class_0"

# Styling constants matching kec template.pptx
FONT_NAME = "Times New Roman"
COLOR_TITLE = RGBColor(192, 0, 0)        # #C00000 Crimson Red / Maroon (KEC Primary)
COLOR_ACCENT = RGBColor(15, 111, 198)    # #0F6FC6 KEC Template Blue
COLOR_BODY = RGBColor(30, 30, 30)        # Dark Charcoal / Black
COLOR_MUTED = RGBColor(110, 110, 110)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_TABLE_HDR = RGBColor(192, 0, 0)
COLOR_ROW_ALT = RGBColor(245, 247, 250)

def set_shape_text(shape, text, font_size=18, bold=False, color=COLOR_BODY, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT_NAME
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return p

def add_header(slide, title_text, category_text=None):
    # Header title box
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.9))
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    if category_text:
        p_cat = tf.paragraphs[0]
        p_cat.alignment = PP_ALIGN.LEFT
        r_cat = p_cat.add_run()
        r_cat.text = category_text.upper()
        r_cat.font.name = FONT_NAME
        r_cat.font.size = Pt(11)
        r_cat.font.bold = True
        r_cat.font.color.rgb = COLOR_ACCENT
        p_title = tf.add_paragraph()
    else:
        p_title = tf.paragraphs[0]
        
    p_title.alignment = PP_ALIGN.LEFT
    r_title = p_title.add_run()
    r_title.text = title_text
    r_title.font.name = FONT_NAME
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_TITLE
    
    # Subtle horizontal line below header
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.28), Inches(11.73), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(220, 220, 220)
    line.line.color.rgb = RGBColor(220, 220, 220)

def add_footer(slide, slide_num, total_slides=21):
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.73), Inches(0.3))
    tf = footer_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    
    r_proj = p.add_run()
    r_proj.text = "Project 27PR27 | Department of Information Technology, Kongu Engineering College"
    r_proj.font.name = FONT_NAME
    r_proj.font.size = Pt(10)
    r_proj.font.color.rgb = COLOR_MUTED
    
    r_space = p.add_run()
    r_space.text = "\t\t\t\t\t\t\t\t\t\t\t\t"
    
    r_num = p.add_run()
    r_num.text = str(slide_num)
    r_num.font.name = FONT_NAME
    r_num.font.size = Pt(10)
    r_num.font.bold = True
    r_num.font.color.rgb = COLOR_TITLE

def add_bullet_list(slide, points, left=0.8, top=1.45, width=11.7, height=5.4, font_size=16):
    tx_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    for i, pt in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.level = 0
        
        if isinstance(pt, tuple):
            title, desc = pt
            r_title = p.add_run()
            r_title.text = "•  " + title + " "
            r_title.font.name = FONT_NAME
            r_title.font.size = Pt(font_size)
            r_title.font.bold = True
            r_title.font.color.rgb = COLOR_BODY
            
            r_desc = p.add_run()
            r_desc.text = desc
            r_desc.font.name = FONT_NAME
            r_desc.font.size = Pt(font_size)
            r_desc.font.color.rgb = COLOR_BODY
        else:
            r = p.add_run()
            r.text = "•  " + pt
            r.font.name = FONT_NAME
            r.font.size = Pt(font_size)
            r.font.color.rgb = COLOR_BODY

def add_table_slide(slide, headers, rows, left=0.8, top=1.5, width=11.73, height=5.2, col_widths=None):
    num_rows = len(rows) + 1
    num_cols = len(headers)
    table_shape = slide.shapes.add_table(num_rows, num_cols, Inches(left), Inches(top), Inches(width), Inches(height))
    table = table_shape.table
    
    if col_widths and len(col_widths) == num_cols:
        for c_idx, w in enumerate(col_widths):
            table.columns[c_idx].width = Inches(w)
            
    # Format Header
    for c_idx, text in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_TABLE_HDR
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.name = FONT_NAME
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = COLOR_WHITE
        
    # Format Rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            if r_idx % 2 == 1:
                cell.fill.fore_color.rgb = COLOR_ROW_ALT
            else:
                cell.fill.fore_color.rgb = COLOR_WHITE
                
            p = cell.text_frame.paragraphs[0]
            val_str = str(val).strip()
            # Align center for short numbers/IDs
            if len(val_str) <= 6 and (val_str.isdigit() or val_str.endswith("%") or val_str in ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10"]):
                p.alignment = PP_ALIGN.CENTER
            else:
                p.alignment = PP_ALIGN.LEFT
                
            run = p.add_run()
            run.text = val_str
            run.font.name = FONT_NAME
            run.font.size = Pt(11)
            run.font.color.rgb = COLOR_BODY

def add_two_column_slide(slide, col1_content, col2_content, left=0.8, top=1.45, width=11.73, height=5.3):
    col_w = (width - 0.4) / 2
    # Left Column
    tx_left = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(col_w), Inches(height))
    tf1 = tx_left.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0
    for i, pt in enumerate(col1_content):
        p = tf1.paragraphs[0] if i == 0 else tf1.add_paragraph()
        p.space_after = Pt(8)
        if isinstance(pt, tuple):
            title, desc = pt
            r_t = p.add_run()
            r_t.text = "• " + title + " "
            r_t.font.name = FONT_NAME
            r_t.font.size = Pt(15)
            r_t.font.bold = True
            r_t.font.color.rgb = COLOR_BODY
            
            r_d = p.add_run()
            r_d.text = desc
            r_d.font.name = FONT_NAME
            r_d.font.size = Pt(15)
            r_d.font.color.rgb = COLOR_BODY
        else:
            r = p.add_run()
            r.text = "• " + pt
            r.font.name = FONT_NAME
            r.font.size = Pt(15)
            r.font.color.rgb = COLOR_BODY
            
    # Right Column
    tx_right = slide.shapes.add_textbox(Inches(left + col_w + 0.4), Inches(top), Inches(col_w), Inches(height))
    tf2 = tx_right.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    for i, pt in enumerate(col2_content):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.space_after = Pt(8)
        if isinstance(pt, tuple):
            title, desc = pt
            r_t = p.add_run()
            r_t.text = "• " + title + " "
            r_t.font.name = FONT_NAME
            r_t.font.size = Pt(15)
            r_t.font.bold = True
            r_t.font.color.rgb = COLOR_BODY
            
            r_d = p.add_run()
            r_d.text = desc
            r_d.font.name = FONT_NAME
            r_d.font.size = Pt(15)
            r_d.font.color.rgb = COLOR_BODY
        else:
            r = p.add_run()
            r.text = "• " + pt
            r.font.name = FONT_NAME
            r.font.size = Pt(15)
            r.font.color.rgb = COLOR_BODY

def add_image_slide(slide, img_path, left=0.8, top=1.45, width=7.0, height=5.2, side_bullets=None, caption=None):
    if Path(img_path).exists():
        slide.shapes.add_picture(str(img_path), Inches(left), Inches(top), width=Inches(width))
    else:
        # Placeholder box
        ph = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        set_shape_text(ph, f"[IMAGE ASSET: {Path(img_path).name} TO BE INSERTED]", font_size=14, bold=True, color=COLOR_TITLE, align=PP_ALIGN.CENTER)
        
    if caption:
        cap_box = slide.shapes.add_textbox(Inches(left), Inches(top + height - 0.4), Inches(width), Inches(0.4))
        set_shape_text(cap_box, caption, font_size=11, bold=True, color=COLOR_MUTED, align=PP_ALIGN.CENTER)
        
    if side_bullets:
        bullet_left = left + width + 0.4
        bullet_w = 11.73 - bullet_left + 0.8
        tx_box = slide.shapes.add_textbox(Inches(bullet_left), Inches(top), Inches(bullet_w), Inches(height))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        for i, pt in enumerate(side_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            if isinstance(pt, tuple):
                title, desc = pt
                r_t = p.add_run()
                r_t.text = "• " + title + ": "
                r_t.font.name = FONT_NAME
                r_t.font.size = Pt(15)
                r_t.font.bold = True
                r_t.font.color.rgb = COLOR_BODY
                
                r_d = p.add_run()
                r_d.text = desc
                r_d.font.name = FONT_NAME
                r_d.font.size = Pt(15)
                r_d.font.color.rgb = COLOR_BODY
            else:
                r = p.add_run()
                r.text = "• " + pt
                r.font.name = FONT_NAME
                r.font.size = Pt(15)
                r.font.color.rgb = COLOR_BODY

print("Helper functions defined successfully.")
