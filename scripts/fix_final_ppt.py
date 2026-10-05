"""
Post-process the final PPTX:
1. Remove stale Thank-You at slide index 6
2. Insert Proposed Methodology slide at correct position (after Problem Description)
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent))
import build_kec_final_ppt as B

OUT = Path("c:/Rice classifier final project/27PR27_KEC_Final_Presentation.pptx")

prs = Presentation(str(OUT))
print(f"Loaded: {len(prs.slides)} slides")
for i, s in enumerate(prs.slides):
    txts = [sh.text_frame.text[:35].replace('\n',' ').strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    imgs = sum(1 for sh in s.shapes if sh.shape_type == 13)
    print(f"  {i}: {txts[0][:40] if txts else '[empty]'}  imgs={imgs}")

# Remove stale slide at index 6 (the old template Thank-You)
sldIdLst = prs.slides._sldIdLst
el6 = sldIdLst[6]
rid6 = el6.get(qn("r:id"))
print(f"\nRemoving stale slide 7 (index 6), rId={rid6}...")
sldIdLst.remove(el6)
try:
    prs.part.drop_rel(rid6)
except Exception as e:
    print(f"  drop_rel: {e}")

# Add Proposed Methodology slide at end first
B.cslide(prs, "Proposed Methodology", 7, B.s_method)
print("Added Proposed Methodology at end")

# Move it to position 6 (after Problem Description at index 5)
sldIdLst = prs.slides._sldIdLst
last_el = sldIdLst[-1]
sldIdLst.remove(last_el)
sldIdLst.insert(6, last_el)
print("Moved to position index 6 (slide 7)")

print(f"\nFinal: {len(prs.slides)} slides")
for i, s in enumerate(prs.slides):
    txts = [sh.text_frame.text[:35].replace('\n',' ').strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    imgs = sum(1 for sh in s.shapes if sh.shape_type == 13)
    print(f"  Slide {i+1:2d}: {txts[0][:42] if txts else '[empty]'}  imgs={imgs}")

prs.save(str(OUT))
print("Saved.")
