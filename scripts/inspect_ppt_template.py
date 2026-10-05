import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = pptx.Presentation("kec template.pptx")
print(f"Total slides in template: {len(prs.slides)}")
print(f"Slide width: {prs.slide_width.inches:.2f} inches, Slide height: {prs.slide_height.inches:.2f} inches")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx+1} ====================")
    title_text = "NO_TITLE"
    if slide.shapes.title and slide.shapes.title.text:
        title_text = slide.shapes.title.text.strip().replace("\n", " ")
    print(f"Slide Title: {title_text}")
    print(f"Layout Name: {slide.slide_layout.name}")
    
    for s_idx, shape in enumerate(slide.shapes):
        shape_type = shape.shape_type
        name = shape.name
        left = shape.left.inches if shape.left else 0
        top = shape.top.inches if shape.top else 0
        width = shape.width.inches if shape.width else 0
        height = shape.height.inches if shape.height else 0
        
        info = f"  Shape {s_idx} [{name}] ({shape_type}) @ ({left:.2f}, {top:.2f}, {width:.2f}x{height:.2f})"
        
        if shape.has_text_frame:
            text = shape.text_frame.text.strip().replace("\n", " // ")
            print(f"{info} -> Text: {text[:150]}")
            # inspect paragraphs & fonts
            for p_idx, p in enumerate(shape.text_frame.paragraphs[:3]):
                p_text = p.text.strip()
                if p_text:
                    font_name = p.font.name if p.font else None
                    font_size = p.font.size.pt if (p.font and p.font.size) else None
                    print(f"      P{p_idx}: font={font_name}, size={font_size} -> {p_text[:80]}")
        elif shape.has_table:
            table = shape.table
            print(f"{info} -> Table: {len(table.rows)} rows x {len(table.columns)} cols")
            for r_idx, row in enumerate(table.rows):
                row_vals = [c.text.strip().replace("\n", " ") for c in row.cells]
                print(f"      Row {r_idx}: {' | '.join(row_vals)}")
        elif shape_type == MSO_SHAPE_TYPE.PICTURE:
            print(f"{info} -> Picture")
        else:
            print(info)
