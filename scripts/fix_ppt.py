"""
Fix the generated PPTX:
 - Remove stale slide 11 (old template Thank-You that was not properly deleted)
 - The Results slide is actually there but hidden by the stale slide
"""
from pptx import Presentation
from pptx.oxml.ns import qn
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).parent))
import build_kec_ppt as B

OUT = Path("c:/Rice classifier final project/27PR27_Rice_Classifier_Presentation.pptx")

prs = Presentation(str(OUT))
print(f"Loaded: {len(prs.slides)} slides")

for i, s in enumerate(prs.slides):
    txts = [sh.text_frame.text[:40].replace('\n', ' ').strip()
            for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    imgs = sum(1 for sh in s.shapes if sh.shape_type == 13)
    first = txts[0] if txts else "[empty]"
    print(f"  Slide {i+1:2d}: {first[:45]!r}  imgs={imgs}")

# Slide 11 (index 10) is the stale Thank-You -- remove it
sldIdLst = prs.slides._sldIdLst
el10 = sldIdLst[10]
rid10 = el10.get(qn("r:id"))
print(f"\nRemoving stale slide 11 (rId={rid10})...")
sldIdLst.remove(el10)
try:
    prs.part.drop_rel(rid10)
except Exception as e:
    print(f"  drop_rel: {e}")

print(f"Slides after: {len(prs.slides)}")
prs.save(str(OUT))
print("Saved OK.\n")

# Final verification
prs2 = Presentation(str(OUT))
print(f"FINAL: {len(prs2.slides)} slides")
for i, s in enumerate(prs2.slides):
    txts = [sh.text_frame.text[:40].replace('\n', ' ').strip()
            for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    imgs = sum(1 for sh in s.shapes if sh.shape_type == 13)
    first = txts[0] if txts else "[empty]"
    print(f"  Slide {i+1:2d}: {first[:45]!r}  imgs={imgs}")
